from __future__ import annotations

import json
import os
import re
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable, Iterator

import nibabel as nib
import numpy as np
from openpyxl import load_workbook

from core.exceptions import InvalidInputError
from data.structures import CompetitionDataset, Series, Study

#: 加载容错开关。默认 **关闭**（保持团队既有的 fail-fast 契约，
#: `tests/test_streaming.py` 对此有明确断言）。
#: 正式评测时由启动脚本置 ``GLIOMA_LOADER_TOLERANT=1``：
#: 单个序列目录/文件异常时"跳过并继续"，而不是让**整批**评测失败
#: （赛事评测不可重跑，1 个脏文件导致全盘 0 分的代价远大于丢 1 例）。
_TOLERANT = os.environ.get("GLIOMA_LOADER_TOLERANT", "").strip().lower() in {"1", "true", "yes", "on"}


_NIFTI_SUFFIXES = (".nii", ".nii.gz")

#: 掩膜文件名关键词（**与训练侧 ``shared/data.py`` 的 ``MASK_HINTS`` 同一语义**）。
#:
#: 掩膜常与影像放在**同一个序列目录**里（本地模拟集即 ``flair_0000/flair.nii.gz``
#: 与 ``flair_0000/瘤体.nii.gz`` 并存）。早期这里只列英文词，于是中文掩膜被当成
#: 第二个序列文件：
#:
#: * 非容错模式 → :func:`_select_original_nifti_files` 抛错，**整批评测失败**；
#: * 容错模式   → 跳过整个序列目录，该模态**整体丢失**并降级推理。
#:
#: 两种后果都与"只是多了个标签文件"不相称，因此这里补齐中文词：
#: 漏检代价是全盘失败/静默降级，误检代价只是少读一个本就该忽略的文件。
_MASK_HINTS = ("mask", "seg", "label", "roi", "掩码", "标注",
               "瘤体", "水肿", "异常", "核心", "病灶", "肿瘤区")

_SERIES_TYPE_HEADERS = ("accessionnumber", "seriesuid", "seriestype")

#: 类型表的列名候选（归一化后做**子串**匹配，故短词靠后）。
#:
#: 为什么不能只认 :data:`_SERIES_TYPE_HEADERS` 里的三个精确名：官方表是**中文/英文
#: 混着来**的（另一份表就叫"检查号 / 序列号 / 序列类型"）。列名对不上时旧实现直接抛
#: ``InvalidInputError`` —— 而评测期**不可重跑**，一次列名改版就是整批 0 分。
_SERIES_TYPE_ALIASES: dict[str, tuple[str, ...]] = {
    "acc": ("accessionnumber", "accession", "检查号", "检查编号", "病例号"),
    "uid": ("seriesinstanceuid", "seriesuid", "序列号", "序列uid"),
    "typ": ("seriestype", "序列类型", "模态", "序列描述", "type"),
}

#: 表**可能**落在的位置（相对数据根）。表与影像同层，但数据根常被指到
#: 阶段目录 / 病例目录 / 容器层的上一层，所以把相邻位置一并试掉。
#: 只认 ``root/SeriesType.xlsx`` 的话，差一层目录就整表读不到。
_TABLE_CANDIDATE_SUBDIRS = ("", "annotation", "original")

#: 表头行扫描上限：第 1~3 行都可能是"索引信息"（标题/说明/空行），
#: 所以在前若干行里找"能凑齐三列名"的那一行，不假设它在第几行。
_HEADER_SCAN_ROWS = 20
#: 按取值嗅探时最多看多少行（够算比例即可，不把全表读进内存）。
_SNIFF_SCAN_ROWS = 300

#: 模态取值的长度上限：``T1CE（增强）`` 也就 8 个字符，检查号 / UID 这类长串必不是模态。
_MODALITY_VALUE_MAXLEN = 16
#: 这些取值本身不是模态（``其他`` 是**权威排除**），但出现在类型列里说明"这列是类型列"。
_OTHER_VALUE_TOKENS = frozenset({
    "其他", "其它", "无", "没有", "other", "none", "na", "n/a", "正常", "平扫",
})

#: 列别名 → 代码里统一使用的规范列名（都归一到 :data:`_SERIES_TYPE_HEADERS`）。
_SERIES_TYPE_CANON = {"acc": "accessionnumber", "uid": "seriesuid", "typ": "seriestype"}

#: 已就"行缺值被跳过"告警过（避免每行刷一条；评测日志要能一眼看到）
_WARNED_PARTIAL_ROW = False

#: 顶层**非病例**目录（与训练侧 ``discover_cases`` 的跳过列表保持同一语义）。
#:
#: 官方数据根除病例号目录外还有标注目录 ``annotation/{fake,Composition,duplicate}``。
#: 不跳过时 ``annotation`` 会被当成一个 accession，后果依次加重：
#:
#: 1. 其子目录里的 NIfTI 会被逐体素读入（白耗算力）；
#: 2. 产出 ``answer/<evaluation_id>/annotation/prediction.json`` 这类垃圾结果，
#:    而平台按检查号评分，多余目录可能被判格式错误；
#: 3. 更致命的是若 ``annotation/fake/<uid>/`` 下的文件名不等于目录名，
#:    :func:`_select_original_nifti_files` 会**直接抛错**——一次评测机会全盘报废。
#:    （赛事评测不可重跑，1 个标注目录不该让整批归零。）
#:
#: ``original`` 是**验证集**的容器层（``verification/original/<检查号>/``）。
#: 它在正常情况下会被 :meth:`DatasetLoader._resolve_dataset_root` 下钻掉；
#: 万一没下钻（dataset_path 直接指 ``…/verification`` 且本层还混了别的目录），
#: 把它列进来至少让"扫不到 NIfTI"响亮失败，而不是拿容器名当检查号产出垃圾答案。
_NON_CASE_DIRS = frozenset({"annotation", "original", "cache", "runs", "folds",
                            "labels"})

#: 平台 ``/2026aicompetition/datasets`` 下的阶段目录名。
#: ``dataset_path`` 若误指**父目录**，这些名字会被当成病例号，静默产出 5 份垃圾答案；
#: 因此显式拦截（或唯一时自动下钻），把"静默全错"变成"当场可见"。
_PLATFORM_PHASES = frozenset({
    "training",
    "evaluation_first",
    "evaluation_second",
    "evaluation_finals",
    "verification",
})

#: 阶段目录**内部**再套的容器层名：影像、``SeriesType.xlsx``、标注表都在它下面。
#: 实测：训练集是 ``training/annotation/``，验证集是 ``verification/original/``。
_CONTAINER_DIRS = frozenset({"annotation", "original"})


def _nifti_stem(path: Path) -> str:
    return path.name[:-7] if path.name.lower().endswith(".nii.gz") else path.stem


def _is_non_case_path(root: Path, path: Path) -> bool:
    """该 NIfTI 是否位于顶层**非病例**目录之下（如 ``annotation/``）。"""
    relative = path.relative_to(root)
    return len(relative.parts) > 1 and relative.parts[0].casefold() in _NON_CASE_DIRS


def _select_original_nifti_files(root: Path, files: Iterable[Path]) -> list[Path]:
    selected: list[Path] = []
    grouped: dict[Path, list[Path]] = defaultdict(list)
    for path in files:
        if len(path.relative_to(root).parts) <= 2:
            selected.append(path)
        else:
            grouped[path.parent].append(path)

    for directory, paths in grouped.items():
        if len(paths) == 1:
            selected.extend(paths)
            continue
        originals = [path for path in paths if _nifti_stem(path) == directory.name]
        if len(originals) != 1:
            if not _TOLERANT:
                raise InvalidInputError(
                    f"series directory {directory} has multiple NIfTI files but "
                    f"expected exactly one original named {directory.name}.nii or "
                    f"{directory.name}.nii.gz: {[path.name for path in sorted(paths)]}"
                )
            print(
                f"[loader][容错] 跳过不合规序列目录 {directory}: "
                f"{[path.name for path in sorted(paths)]}",
                flush=True,
            )
            continue
        selected.extend(originals)
    return sorted(selected)


def _clean_identifier(value: Any, fallback: str) -> str:
    text = str(value or "").strip()
    if not text:
        text = fallback
    return re.sub(r"[\\/\x00-\x1f]", "_", text)


def _metadata_key(value: Any) -> str:
    return re.sub(r"\s+", "", str(value if value is not None else "")).casefold()


def _looks_like_modality_value(value: Any) -> bool:
    """该单元格"看着像模态取值"吗（专供列名认不出时的取值嗅探）。

    判据与模态匹配**同源**（``data.series_selector.guess_modality``，局部导入）：
    嗅探认定"这列是类型列"的取值，后面也得真能被认出来 —— 两套口径不一致，
    会出现"列挑对了、模态仍全空"这种最难查的情形。
    """
    text = re.sub(r"[\s\-_/]+", "", str(value if value is not None else "")).casefold()
    if not text or len(text) > _MODALITY_VALUE_MAXLEN:
        return False
    if text in _OTHER_VALUE_TOKENS:
        return True
    from data.series_selector import guess_modality           # 局部导入：不引入模块级耦合
    return guess_modality(text) is not None


def _match_header_row(row) -> dict[str, int] | None:
    """这一行是不是类型表的表头（三列都能按别名认出来）→ ``{规范列名: 下标}``。"""
    cells = [_metadata_key(value) for value in row]
    if not any(cells):
        return None
    found: dict[str, int] = {}
    for want, keys in _SERIES_TYPE_ALIASES.items():
        for index, cell in enumerate(cells):
            if cell and any(k in cell for k in keys):
                found[_SERIES_TYPE_CANON[want]] = index
                break
    return found if set(found) == set(_SERIES_TYPE_HEADERS) else None


def _sniff_series_type_columns(rows: list) -> tuple[int, int, int] | None:
    """**不看列名**，按取值找出 ``(检查号列, 序列号列, 类型列)``；认不出返回 ``None``。

    为什么需要它：列名是唯一会被"改版"的东西（前两列写成 ``编号/影像号``、
    加了索引列、或干脆是 ``A/B/C``），而**取值**不会变（检查号、DICOM UID、
    5 类模态取值）。评测期不可重跑，多这一层兜底就少一种整批 0 分的方式。

    判据（全在取值上）：类型列 = "像模态取值"的比例最高且 ≥ 0.5；
    剩下两列里取值含 ``.``（DICOM UID）比例更高 / 平均更长的那个是序列号列。

    找不到就返回 ``None``（随后仍按原逻辑响亮报错），绝不硬凑 —— 凑错会把整表挂在
    错误的键上，比读不到更难查。
    """
    body = [r for r in rows[:_SNIFF_SCAN_ROWS]
            if any(str(c).strip() for c in r if c is not None)]
    if len(body) < 3:
        return None
    width = max(len(r) for r in body)
    if width < 3:
        return None
    cols = [[str(r[c]).strip() for r in body if c < len(r) and str(r[c]).strip()]
            for c in range(width)]
    ratio = [(sum(1 for v in vals if _looks_like_modality_value(v)) / len(vals)
              if vals else 0.0) for vals in cols]
    typ_col = max(range(width), key=lambda c: (ratio[c], len(cols[c])))
    if ratio[typ_col] < 0.5:
        return None
    rest = [c for c in range(width) if c != typ_col and cols[c]]
    if len(rest) < 2:
        return None

    def dotted(c: int) -> float:
        return sum(1 for v in cols[c] if "." in v) / len(cols[c])

    def mean_len(c: int) -> float:
        return sum(len(v) for v in cols[c]) / len(cols[c])

    rest.sort(key=lambda c: (dotted(c), mean_len(c)), reverse=True)
    uid_col, acc_col = rest[0], rest[1]
    return acc_col, uid_col, typ_col


def _series_type_table_path(root: Path) -> Path | None:
    """在数据根本层与相邻层找 ``SeriesType.xlsx`` → 路径（都没有返回 ``None``）。"""
    seen: set[Path] = set()
    for base in (root, root.parent):
        for sub in _TABLE_CANDIDATE_SUBDIRS:
            candidate = (base / sub / "SeriesType.xlsx") if sub else (base / "SeriesType.xlsx")
            if candidate in seen:
                continue
            seen.add(candidate)
            if candidate.is_file():
                return candidate
    return None


def _read_series_types(root: Path) -> dict[tuple[str, str], str]:
    global _WARNED_PARTIAL_ROW
    path = _series_type_table_path(root)
    if path is None:
        print(f"[loader][告警] {root} 及其相邻层都没找到 SeriesType.xlsx → "
              f"序列类型读不到（将退回 sidecar / 文件名关键词 / 兜底表）。", flush=True)
        return {}
    if path.parent != root:
        print(f"[loader][告警] SeriesType.xlsx 不在数据根那一层，已定位到 {path}", flush=True)

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise InvalidInputError(
            f"cannot read series metadata {path}: {type(exc).__name__}: {exc}"
        ) from exc

    try:
        # ① 按**列名**扫表头行：哪一行能凑齐三列就用哪一行，不假设它在第几行
        #    （第 1~3 行都可能是索引信息）。列名走别名子串匹配，中英文表头都认。
        header: tuple[Any, int, dict[str, int]] | None = None
        for worksheet in workbook.worksheets:
            for row_number, row in enumerate(
                worksheet.iter_rows(max_row=_HEADER_SCAN_ROWS, values_only=True),
                start=1,
            ):
                columns = _match_header_row(row)
                if columns is not None:
                    header = (worksheet, row_number, columns)
                    break
            if header is not None:
                break

        if header is None:
            # ② 列名一条都没命中（改版 / 加了索引列 / 表头用了别的词）
            #    → 按**取值**嗅探三列。嗅探也失败才响亮报错（原行为）。
            for worksheet in workbook.worksheets:
                sample = list(worksheet.iter_rows(max_row=_SNIFF_SCAN_ROWS,
                                                  values_only=True))
                sniff = _sniff_series_type_columns(sample)
                if sniff is None:
                    continue
                acc_col, uid_col, typ_col = sniff
                print(f"[loader][告警] {path.name} 工作表 {worksheet.title!r} 列名未识别 → "
                      f"已按取值定位：检查号=第 {acc_col + 1} 列、序列号=第 {uid_col + 1} 列、"
                      f"类型=第 {typ_col + 1} 列。若取值明显不对，请把该表前几行贴出来。",
                      flush=True)
                header = (worksheet, 0, {"accessionnumber": acc_col,
                                         "seriesuid": uid_col,
                                         "seriestype": typ_col})
                break

        if header is None:
            raise InvalidInputError(
                f"series metadata {path} is missing headers: "
                "AccessionNumber / SeriesUid / SeriesType"
                f"（已扫前 {_HEADER_SCAN_ROWS} 行，并尝试按取值嗅探）"
            )

        worksheet, header_row, columns = header
        sniffed = header_row == 0
        series_types: dict[tuple[str, str], str] = {}
        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            values = {
                name: row[index] if index < len(row) else None
                for name, index in columns.items()
            }
            if all(
                not str(value if value is not None else "").strip()
                for value in values.values()
            ):
                continue
            if sniffed and not _looks_like_modality_value(values["seriestype"]):
                continue        # 嗅探模式：表头行 / 说明行的"类型"取值不像模态，滤掉
            missing = [
                name
                for name, value in values.items()
                if not str(value if value is not None else "").strip()
            ]
            if missing:
                # 表尾的合计/备注行常常只填一两列。旧实现直接抛错 → **整批评测失败**，
                # 与"一行残缺"的代价完全不相称；跳过并告警即可（同一键冲突仍会响亮报错）。
                if not _WARNED_PARTIAL_ROW:
                    _WARNED_PARTIAL_ROW = True
                    print(f"[loader][告警] {path.name} 工作表 {worksheet.title!r} 第 "
                          f"{row_number} 行缺 {', '.join(missing)} → 已跳过该行"
                          f"（表尾备注行常见；本进程内只提示这一次）。", flush=True)
                continue

            key = (
                _metadata_key(values["accessionnumber"]),
                _metadata_key(values["seriesuid"]),
            )
            series_type = str(values["seriestype"]).strip()
            previous = series_types.get(key)
            if previous is not None and previous != series_type:
                raise InvalidInputError(
                    f"series metadata {path} has conflicting SeriesType values "
                    f"for accession={values['accessionnumber']!r} "
                    f"series={values['seriesuid']!r}: {previous!r}, {series_type!r}"
                )
            series_types[key] = series_type
        return series_types
    except InvalidInputError:
        raise
    except Exception as exc:
        raise InvalidInputError(
            f"cannot read series metadata {path}: {type(exc).__name__}: {exc}"
        ) from exc
    finally:
        workbook.close()


class DatasetLoader:
    """Load a NIfTI tree into the competition domain model."""

    def load(self, dataset_path: str | Path) -> CompetitionDataset:
        """Compatibility entry point for callers that need the full dataset."""
        root = Path(dataset_path).expanduser().resolve()
        return CompetitionDataset(
            root,
            tuple(self.iter_studies(root)),
            {"format": "nifti"},
        )

    def iter_studies(self, dataset_path: str | Path) -> Iterator[Study]:
        """Discover file paths, then load one study at a time."""
        root = self._resolve_dataset_root(dataset_path)
        if not root.is_dir():
            raise InvalidInputError(f"dataset_path is not a directory: {root}")

        nifti_files = _select_original_nifti_files(
            root,
            (
                path
                for path in root.rglob("*")
                if path.is_file()
                and path.name.lower().endswith(_NIFTI_SUFFIXES)
                and not any(hint in path.name.lower() for hint in _MASK_HINTS)
                and not _is_non_case_path(root, path)
            ),
        )
        if not nifti_files:
            raise InvalidInputError(f"no readable NIfTI images under {root}")
        yield from self._iter_nifti(root, nifti_files, _read_series_types(root))

    @staticmethod
    def _resolve_dataset_root(dataset_path: str | Path) -> Path:
        """确认 ``dataset_path`` 落在**含病例号目录的那一层**。

        平台的 ``/2026aicompetition/datasets`` 下有 5 个阶段目录
        （``training`` / ``evaluation_first`` / ``evaluation_second`` /
        ``evaluation_finals`` / ``verification``），而 ``dataset_path`` 必须精确到
        其中**一个**。若误传父目录，旧实现会把阶段名当成病例号：

        * 产出 5 个以阶段命名的"检查"，与真实检查集完全对不上；
        * 平台按检查号评分 → 全部缺失，却不报任何错。

        这里唯一能安全补救的情形是"父目录下只有一个阶段目录"（自动下钻并告警）；
        真正有多个候选时**必须报错**——猜错阶段会把答案写到错误的评测轮次上，
        比直接失败更糟。
        """
        root = Path(dataset_path).expanduser().resolve()
        if not root.is_dir():
            return root

        children = sorted(p.name for p in root.iterdir() if p.is_dir())
        if not children or not {name.casefold() for name in children} <= _PLATFORM_PHASES:
            # 本层不是"阶段目录的父目录"，但可能是"影像/表的容器层"：
            # ``verification/`` 下只有 ``original/``、``training/`` 下只有 ``annotation/``。
            # 唯一候选时下钻并告警；候选多于一个则维持原样（由后续"扫不到 NIfTI"报错接手，
            # 猜错层比直接失败更糟）。
            if len(children) == 1 and children[0].casefold() in _CONTAINER_DIRS:
                print(f"[loader][告警] dataset_path={root} 下没有检查号目录，"
                      f"已自动下钻到 {children[0]}/", flush=True)
                return root / children[0]
            return root
        if len(children) == 1:
            print(f"[loader][告警] dataset_path={root} 是数据集父目录，"
                  f"已自动下钻到唯一阶段目录 {children[0]}", flush=True)
            return root / children[0]
        raise InvalidInputError(
            f"dataset_path 指向数据集父目录 {root}，其下是平台阶段目录 {children}。"
            f"请指向**具体阶段**，例如 {root / 'evaluation_first'}；"
            f"否则阶段名会被当作病例号，答案目录将整体错位。"
        )

    def _iter_nifti(
        self,
        root: Path,
        files: Iterable[Path],
        series_types: dict[tuple[str, str], str],
    ) -> Iterator[Study]:
        grouped: dict[str, list[Path]] = defaultdict(list)
        for path in files:
            relative = path.relative_to(root)
            accession = relative.parts[0] if len(relative.parts) > 1 else _nifti_stem(path)
            grouped[_clean_identifier(accession, "study")].append(path)

        for accession, paths in sorted(grouped.items()):
            series: list[Series] = []
            for path in paths:
                try:
                    series.append(
                        self._read_nifti_series(root, accession, path, series_types)
                    )
                except Exception as exc:                          # noqa: BLE001
                    if not _TOLERANT:
                        raise
                    # 单个序列（文件损坏 / 非 3D / 4D / affine 非法…）读取失败时
                    # 只跳过该序列，不让整个 evaluation 中断。
                    print(
                        f"[loader][容错] 跳过不可读序列 study={accession!r} "
                        f"path={path}: {exc}",
                        flush=True,
                    )
            if series:
                yield Study(accession_number=accession, series=tuple(series))
            else:
                # 该 study 没有任何可用序列 → 不 yield（Study 要求 series 非空）。
                print(
                    f"[loader]{'[容错] ' if _TOLERANT else ' '}study {accession!r} "
                    f"无可用序列",
                    flush=True,
                )

    def _read_nifti_series(
        self,
        root: Path,
        accession: str,
        path: Path,
        series_types: dict[tuple[str, str], str],
    ) -> Series:
        relative = path.relative_to(root)
        if len(relative.parts) == 1 or path.parent == root / relative.parts[0]:
            series_uid = _nifti_stem(path)
        else:
            series_uid = path.parent.name
        sidecar = path.with_name(_nifti_stem(path) + ".json")
        try:
            metadata: dict[str, Any] = {}
            if sidecar.is_file():
                metadata = json.loads(sidecar.read_text(encoding="utf-8"))

            image = nib.load(str(path))
            array = np.asanyarray(image.dataobj)
            array = np.squeeze(array)
            if array.ndim != 3:
                raise InvalidInputError(f"NIfTI image must be 3-D: {path} -> {array.shape}")

            uid = _clean_identifier(
                metadata.get("SeriesInstanceUID"),
                series_uid,
            )
            description = str(
                series_types.get((_metadata_key(accession), _metadata_key(uid)))
                or metadata.get("SeriesDescription")
                or metadata.get("ProtocolName")
                or series_uid
            )
            return Series(
                series_uid=uid,
                modality=description,
                image=np.asarray(array),
                affine=np.asarray(image.affine, dtype=np.float64),
                source_path=path,
                metadata=metadata,
            )
        except Exception as exc:
            detail = (
                str(exc)
                if isinstance(exc, InvalidInputError)
                else f"{type(exc).__name__}: {exc}"
            )
            raise InvalidInputError(
                f"cannot load NIfTI study={accession!r} series={series_uid!r} "
                f"path={path}: {detail}"
            ) from exc
