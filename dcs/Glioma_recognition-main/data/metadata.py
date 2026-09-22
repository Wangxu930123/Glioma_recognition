"""元数据解析：``SeriesType.xlsx`` 与 NIfTI sidecar（规范 §5、§26）。

规范要求把这两项从 ``data/loader.py`` 独立出来——它们只是"**补充**序列类型"的
旁路信息源，与"发现并读取 NIfTI"是两件事，混在一起会让 Loader 难以单测。

解析原则（规范 §6.2）：

- ``SeriesType.xlsx`` 按 ``AccessionNumber + SeriesUid`` 匹配，
  **只补充序列类型、不改写 UID**；
- sidecar JSON 可补充 ``SeriesInstanceUID`` / ``SeriesDescription`` / ``ProtocolName``；
- **解析失败必须携带 accession、series UID 与文件路径**——否则排障时无法定位坏数据。
"""
from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any


def _key(value: Any) -> str:
    """归一化匹配键：去空白、转小写。"""
    return re.sub(r"\s+", "", str(value if value is not None else "")).casefold()


def read_series_types(root: Path) -> dict[tuple[str, str], str]:
    """读取 ``SeriesType.xlsx`` → ``{(accession, series_uid): series_type}``。

    缺文件返回空表（**不是**错误，规范允许回退 sidecar/UID 猜测）；
    同一键出现冲突取值时明确失败（规范 §21 风险表要求"冲突直接失败"）。
    """
    path = Path(root) / "SeriesType.xlsx"
    if not path.is_file():
        return {}
    try:
        from openpyxl import load_workbook
    except ImportError:                                           # pragma: no cover
        return {}

    wb = load_workbook(path, read_only=True, data_only=True)
    out: dict[tuple[str, str], str] = {}
    aliases = {
        "acc": ("accessionnumber", "accession", "检查号", "检查编号"),
        "uid": ("seriesinstanceuid", "seriesuid", "序列号", "序列uid"),
        "typ": ("seriestype", "type", "序列类型", "模态", "序列描述"),
    }
    for sheet in wb.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        header = [_key(c) for c in rows[0]]
        idx: dict[str, int] = {}
        for want, keys in aliases.items():
            for i, h in enumerate(header):
                if any(k in h for k in keys):
                    idx[want] = i
                    break
        if not {"acc", "uid", "typ"} <= idx.keys():
            continue                                              # 该 sheet 不是映射表
        for row in rows[1:]:
            try:
                acc, uid, typ = row[idx["acc"]], row[idx["uid"]], row[idx["typ"]]
            except IndexError:
                continue
            if acc is None or uid is None or typ is None:
                continue
            k = (_key(acc), _key(uid))
            v = str(typ).strip()
            if k in out and out[k] != v:
                raise ValueError(
                    f"SeriesType.xlsx 冲突：accession={acc!r} series={uid!r} "
                    f"同时映射到 {out[k]!r} 与 {v!r}（{path}）")
            out[k] = v
    return out


def read_sidecar(nifti_path: Path) -> dict[str, Any]:
    """读取同名 JSON sidecar；解析失败时**带上完整定位信息**抛出。"""
    p = Path(nifti_path)
    stem = p.name[:-len(".nii.gz")] if p.name.endswith(".nii.gz") else p.stem
    sidecar = p.with_name(stem + ".json")
    if not sidecar.is_file():
        return {}
    try:
        data = json.loads(sidecar.read_text(encoding="utf-8"))
    except Exception as exc:                                      # noqa: BLE001
        raise ValueError(
            f"sidecar 解析失败：accession={p.parent.parent.name!r} "
            f"series={p.parent.name!r} file={sidecar} ({exc})") from exc
    return data if isinstance(data, dict) else {}


def describe_modality(metadata: dict[str, Any], series_types: dict,
                      accession: str, series_uid: str) -> str | None:
    """按规范优先级给出序列类型：xlsx → sidecar → None（由上层按 UID 猜）。"""
    hit = series_types.get((_key(accession), _key(series_uid)))
    if hit:
        return hit
    for k in ("SeriesDescription", "ProtocolName", "SequenceName"):
        v = (metadata or {}).get(k)
        if v:
            return str(v)
    return None
