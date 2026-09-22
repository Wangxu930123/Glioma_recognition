"""提取三份平台指南的「文本 + 图文对应关系」与视频关键帧。

docx 里的图片是**按段落顺序内嵌**的，因此可以还原"哪一步配哪张截图"，
这样才能写出"该点哪里"的指南，而不是把截图堆在一起。
"""
from __future__ import annotations

import os
import pathlib
import subprocess
import zipfile

import docx
from docx.oxml.ns import qn

# 路径自适应：脚本位于 <仓库>/glioma_track4/scripts/，三份官方文档在仓库根
_HERE = pathlib.Path(__file__).resolve().parent
_ROOT = _HERE.parent.parent
SRC = pathlib.Path(os.environ.get("GLIOMA_PLATFORM_DOCS") or _ROOT)
OUT = _HERE.parent / "docs" / ".platform_raw"      # 中间产物（已 gitignore）
OUT.mkdir(parents=True, exist_ok=True)

DOCS = ["云电脑平台使用指南V1.0", "训推平台使用指南V1.0", "赛事管理平台使用指南V2.0"]


def dump(doc_path: pathlib.Path, tag: str) -> None:
    d = docx.Document(str(doc_path))
    # 关系 id → 原始图片名
    rels = {rid: rel.target_ref for rid, rel in d.part.rels.items()
            if "image" in rel.reltype}

    # 1) 按 body 顺序遍历（段落 / 表格），段落内按出现顺序取图片
    lines: list[str] = []
    idx = 0
    for child in d.element.body.iterchildren():
        if child.tag == qn("w:tbl"):
            # 表格：转成 markdown 风格
            tbl = next((t for t in d.tables if t._tbl is child), None)
            if tbl is not None:
                lines.append("\n[表格]")
                for row in tbl.rows:
                    lines.append("  | " + " | ".join(c.text.strip() for c in row.cells))
                lines.append("")
            continue
        if child.tag != qn("w:p"):
            continue
        para = next((p for p in d.paragraphs if p._p is child), None)
        if para is None:
            continue
        text = para.text.strip()
        style = para.style.name if para.style is not None else ""
        imgs: list[str] = []
        for blip in child.findall(".//" + qn("a:blip")):
            rid = blip.get(qn("r:embed"))
            if rid and rid in rels:
                idx += 1
                imgs.append(f"⟦IMG{idx}:{pathlib.Path(rels[rid]).name}⟧")
        if text or imgs:
            prefix = "## " if style.startswith(("Heading", "标题")) else ""
            lines.append(f"{prefix}{text} {' '.join(imgs)}".rstrip())

    (OUT / f"{tag}.txt").write_text("\n".join(lines), encoding="utf-8")

    # 2) 抽出图片文件
    img_dir = OUT / f"{tag}_imgs"
    img_dir.mkdir(exist_ok=True)
    with zipfile.ZipFile(doc_path) as z:
        for name in z.namelist():
            if name.startswith("word/media/"):
                data = z.read(name)
                (img_dir / pathlib.Path(name).name).write_bytes(data)
    n_img = len(list(img_dir.iterdir()))
    print(f"  {tag}: 文本 {len(lines)} 行，图片 {n_img} 张 → {img_dir}")


print("=== 提取三份指南 ===")
for name in DOCS:
    dump(SRC / f"{name}.docx", name)

# ------------------------------------------------------------------ #
print("\n=== 视频信息与关键帧 ===")
mp4 = SRC / "平台使用方法.mp4"
probe = subprocess.run(
    ["ffprobe", "-v", "error", "-show_entries",
     "format=duration,size:stream=width,height,r_frame_rate",
     "-of", "default=nw=1", str(mp4)],
    capture_output=True, text=True)
print(probe.stdout.strip())

frames = OUT / "frames"
frames.mkdir(exist_ok=True)
subprocess.run(["ffmpeg", "-v", "error", "-y", "-i", str(mp4),
                "-vf", "fps=1/8,scale=1024:-1", "-q:v", "4",
                str(frames / "f%03d.jpg")], check=False)
got = sorted(frames.glob("*.jpg"))
print(f"  提取关键帧 {len(got)} 张（每 8 秒一张）→ {frames}")
